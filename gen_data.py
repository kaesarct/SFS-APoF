import openpyxl, json, os

base = os.path.dirname(os.path.abspath(__file__))

wb_ita = openpyxl.load_workbook(os.path.join(base, 'doc', 'Elenco_Paesi_in_italiano.xlsx'))

# --- ita-to-eng map (col A = ita names in english order, zip with english names) ---
ws_ita = wb_ita['Elenco_Paesi_ITA']
rows_ita = list(ws_ita.iter_rows(min_row=1, values_only=True))
col_a = [r[0] for r in rows_ita if r[0]]

wb_eng = openpyxl.load_workbook(os.path.join(base, 'doc', 'Elenco_Paesi.xlsx'))
ws_eng = wb_eng.active
eng_names = [r[0] for r in ws_eng.iter_rows(min_row=1, values_only=True) if r[0]]

ita_to_eng = dict(zip(col_a, eng_names))
out1 = os.path.join(base, 'src', 'app', 'ita-to-eng.json')
with open(out1, 'w', encoding='utf-8') as f:
    json.dump(ita_to_eng, f, ensure_ascii=False, indent=2)
print(f'ita-to-eng: {len(ita_to_eng)} entries -> {out1}')

# --- paesi-ita (col E = alphabetical italian order) ---
col_e = [r[4] for r in rows_ita if r[4] and r[4] != 'In ordine alfabetico per menu a tendina']
out2 = os.path.join(base, 'src', 'app', 'paesi-ita.json')
with open(out2, 'w', encoding='utf-8') as f:
    json.dump(col_e, f, ensure_ascii=False, indent=2)
print(f'paesi-ita: {len(col_e)} entries -> {out2}')

# --- currency-data (from Valute sheet) ---
ws_val = wb_ita['Valute']
rows_val = list(ws_val.iter_rows(min_row=2, values_only=True))
currency_data = {r[1]: {'currency': r[2], 'rate': r[3]} for r in rows_val if r[1] and r[3]}
out3 = os.path.join(base, 'src', 'app', 'currency-data.json')
with open(out3, 'w', encoding='utf-8') as f:
    json.dump(currency_data, f, ensure_ascii=False, indent=2)
print(f'currency-data: {len(currency_data)} entries -> {out3}')

print('DONE')
